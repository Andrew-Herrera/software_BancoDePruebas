# storage/file_manager.py
import os
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class FileManager:
    """
    Capa de Almacenamiento: Gestor exclusivo de archivos (JSON, CSV, XLSX).
    Agnóstico de la base de datos y de la lógica de la interfaz.
    """

    def __init__(self, data_dir: str = "data"):
        self.data_dir = data_dir
        self.config_path = os.path.join(self.data_dir, "config.json")
        
        # Asegurar que el directorio de datos exista
        if not os.path.exists(self.data_dir):
            os.makedirs(self.data_dir)

    @staticmethod
    def get_default_config() -> dict:
        """Fuente de verdad para los valores por defecto del sistema."""
        return {
            # Valores de prueba por defecto
            "v_start": -0.200,
            "v_stop": 1.200,
            "samples": 100,
            "delay": 0.05,
            "compliance": 0.05,
            "nplc": 1.0,
            "hysteresis": False,
            "last_irr": 1000.0,
            "n_sweeps": 1,
            "int_sweeps": 5.0,
            "dur_hrs": 1.0,
            "b_per_hr": 10,
            "v_bias": 0.0,
            
            # Preferencias de conexión y hardware
            "demo_mode": True,
            "gpib_id": "10",
            "baud": "9600",
            "termination": "\\r\\n",
            "term_route": "Front (FRONT)",
            "wire_mode": "2 Wire (LOCAL)",
            
            # Metrología y UI
            "sensor": "BPW20RF",
            "last_calib": "2026-04-09",
            "remind_days": 10,
            "graph_polarity": "Direct (PV - Quadrant I)",
            "strict_validation": True,
            "lang": "Español",
            "dark_mode": True,
            "touch_mode": True
        }

    def load_json_config(self) -> dict:
        """Carga el JSON de preferencias. Si no existe o falla, crea uno con los defaults."""
        defaults = self.get_default_config()

        if not os.path.exists(self.config_path):
            self.save_json_config(defaults)
            return defaults

        try:
            with open(self.config_path, "r", encoding="utf-8") as f:
                loaded_config = json.load(f)
                # Actualiza los valores por defecto con los que existan en el archivo
                defaults.update(loaded_config)
                return defaults
        except (json.JSONDecodeError, OSError) as e:
            print(f"Error leyendo config.json: {e}. Restaurando valores por defecto.")
            self.save_json_config(defaults)
            return defaults

    def save_json_config(self, config_dict: dict) -> bool:
        """Sobrescribe el archivo de configuración con las preferencias actuales."""
        try:
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(config_dict, f, indent=4, ensure_ascii=False)
            return True
        except OSError as e:
            print(f"Error guardando config.json: {e}")
            return False

    def write_csv_report(self, filepath: str, metadata: dict, curves_points: dict) -> bool:
        """
        Escribe los datos crudos en formato CSV.
        :param curves_points: Diccionario con formato { "S1": [(V, I), ...], "S2": [...] }
        """
        try:
            with open(filepath, mode='w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Bloque de Metadatos
                writer.writerow(["--- METADATOS DEL EXPERIMENTO ---"])
                writer.writerow(["Sustrato (Holder)", metadata.get("holder_id", "")])
                writer.writerow(["Posición (Celda)", metadata.get("posicion", "")])
                writer.writerow(["Autor", metadata.get("autor", "")])
                writer.writerow(["Fecha", metadata.get("fecha", "")])
                writer.writerow(["Notas", metadata.get("notas", "")])
                writer.writerow([])
                
                # Encabezados de columnas de datos
                headers = []
                for sweep_name in curves_points.keys():
                    headers.extend([f"Voltaje_{sweep_name} (V)", f"Corriente_{sweep_name} (A)"])
                writer.writerow(headers)
                
                # Escritura de puntos (alineados por columnas)
                max_pts = max([len(pts) for pts in curves_points.values()]) if curves_points else 0
                for i in range(max_pts):
                    row = []
                    for sweep_name in curves_points.keys():
                        pts = curves_points[sweep_name]
                        if i < len(pts):
                            row.extend([pts[i][0], pts[i][1]])
                        else:
                            row.extend(["", ""])
                    writer.writerow(row)
            return True
        except Exception as e:
            print(f"Error exportando CSV: {e}")
            return False

    def write_excel_report(self, filepath: str, metadata: dict, curves_points: dict, curves_metrics: dict) -> bool:
        """
        Genera un libro Excel con la plantilla institucional, hoja de resumen y datos crudos.
        :param curves_metrics: Diccionario con métricas de mérito { "S1": {"voc": X, "jsc": Y, ...} }
        """
        try:
            wb = Workbook()
            
            # --- HOJA 1: RESUMEN DE PARÁMETROS ---
            ws_summary = wb.active
            ws_summary.title = "Resumen de Parámetros"
            ws_summary.views.sheetView[0].showGridLines = False
            
            # Definición de Estilos
            fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            fill_accent = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            font_title = Font(name="Segoe UI", size=16, bold=True, color="1F497D")
            font_subtitle = Font(name="Segoe UI", size=11, italic=True, color="595959")
            font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_bold = Font(name="Segoe UI", size=10, bold=True)
            font_regular = Font(name="Segoe UI", size=10)
            
            border_thin = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
                                 top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))
            align_center = Alignment(horizontal='center', vertical='center')
            align_left = Alignment(horizontal='left', vertical='center')

            # Encabezado Institucional
            ws_summary.cell(row=2, column=2, value="Reporte de Caracterización J-V de Celdas Solares").font = font_title
            ws_summary.cell(row=3, column=2, value="Laboratorio de Dispositivos Orgánicos - UPIITA IPN").font = font_subtitle
            
            # Metadatos
            metadata_rows = [
                ("Identificador de Sustrato (Holder):", metadata.get("holder_id", "")),
                ("Posición de la Celda:", metadata.get("posicion", "")),
                ("Área Activa [cm²]:", metadata.get("area_celda", 0.09)),
                ("Operador / Investigador:", metadata.get("autor", "")),
                ("Fecha de Extracción:", metadata.get("fecha", "")),
                ("Notas de Muestra:", metadata.get("notas", ""))
            ]
            
            for row_idx, (lbl, val) in enumerate(metadata_rows, start=5):
                c_lbl = ws_summary.cell(row=row_idx, column=2, value=lbl)
                c_lbl.font = font_bold
                c_lbl.alignment = align_left
                c_val = ws_summary.cell(row=row_idx, column=3, value=val)
                c_val.font = font_regular
                c_val.alignment = align_left

            # Tabla de Parámetros Fotovoltaicos
            headers_tabla = ["Barrido", "Voc [V]", "Isc [A]", "Jsc [mA/cm²]", "FF [%]", "Pmax [W]", "Eficiencia [%]"]
            start_row_table = 13
            
            for col_idx, text in enumerate(headers_tabla, start=2):
                cell = ws_summary.cell(row=start_row_table, column=col_idx, value=text)
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
                cell.border = border_thin

            current_row = start_row_table + 1
            for sweep_name, metrics in curves_metrics.items():
                vals = [
                    sweep_name, 
                    metrics.get("voc", 0.0), 
                    metrics.get("isc", 0.0), 
                    metrics.get("jsc", 0.0), 
                    metrics.get("ff", 0.0),
                    metrics.get("pmax", 0.0),
                    metrics.get("eta", 0.0)
                ]
                for col_idx, val in enumerate(vals, start=2):
                    c = ws_summary.cell(row=current_row, column=col_idx, value=val)
                    c.font = font_regular
                    c.border = border_thin
                    c.alignment = align_center
                    if col_idx > 2 and isinstance(val, (int, float)):
                        c.number_format = '0.000E+00' if col_idx == 4 else '0.00'
                current_row += 1

            # --- HOJA 2: DATOS CRUDOS ---
            ws_data = wb.create_sheet(title="Datos Crudos (I-V)")
            ws_data.views.sheetView[0].showGridLines = True
            
            for col_pos, sweep_name in enumerate(curves_points.keys()):
                c_v = ws_data.cell(row=1, column=(col_pos * 2) + 1, value=f"Voltaje_{sweep_name} [V]")
                c_i = ws_data.cell(row=1, column=(col_pos * 2) + 2, value=f"Corriente_{sweep_name} [A]")
                
                for cell in (c_v, c_i):
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="1F497D")
                    cell.fill = fill_accent
                    cell.alignment = align_center
                    cell.border = border_thin

            max_pts = max([len(pts) for pts in curves_points.values()]) if curves_points else 0
            for r_idx in range(max_pts):
                for col_pos, sweep_name in enumerate(curves_points.keys()):
                    pts = curves_points[sweep_name]
                    if r_idx < len(pts):
                        cell_v = ws_data.cell(row=r_idx + 2, column=(col_pos * 2) + 1, value=pts[r_idx][0])
                        cell_i = ws_data.cell(row=r_idx + 2, column=(col_pos * 2) + 2, value=pts[r_idx][1])
                        
                        cell_v.number_format = '0.000'
                        cell_i.number_format = '0.000E+00'
                        
                        for cell in (cell_v, cell_i):
                            cell.font = font_regular
                            cell.border = border_thin
                            cell.alignment = align_center

            # Autoajuste de columnas
            for ws in (ws_summary, ws_data):
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 15)
                    
            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exportando XLSX: {e}")
            return False